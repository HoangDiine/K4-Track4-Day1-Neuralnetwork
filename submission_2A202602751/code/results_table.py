from __future__ import annotations

import json
import math
from pathlib import Path

from openpyxl import load_workbook


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Save configuration, history, and summary without serializing model weights."""
    if not result.get("cfg", {}).get("exp_id"):
        raise ValueError("result.cfg.exp_id is required")
    path = Path(results_dir) / f"{result['cfg']['exp_id']}.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {key: result[key] for key in ("cfg", "history", "summary")}
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False, default=_json_default), encoding="utf-8")
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Load result JSON files, sorted by experiment id."""
    directory = Path(results_dir)
    if not directory.exists():
        return []
    values = [json.loads(p.read_text(encoding="utf-8")) for p in directory.glob("*.json")]
    return sorted(values, key=lambda item: item.get("cfg", {}).get("exp_id", ""))


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Convert a run into one experiments-table row."""
    cfg, summary = result["cfg"], result["summary"]
    row = {**cfg, **summary}
    row["hidden"] = str(tuple(cfg.get("hidden", ())))
    row["figure_file"] = f"figures/{cfg['exp_id']}.png"
    row["notes"] = notes
    if eval_scores:
        row["eval_acc"] = eval_scores.get("accuracy", eval_scores.get("eval_acc"))
        row["eval_macro_f1"] = eval_scores.get("macro_f1", eval_scores.get("eval_macro_f1"))
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Write rows into the template while preserving formula columns and sheets."""
    wb = load_workbook(template_path)
    if "Experiments" not in wb.sheetnames:
        raise KeyError("template is missing the Experiments sheet")
    ws = wb["Experiments"]
    headers = {ws.cell(1, col).value: col for col in range(1, ws.max_column + 1)}
    formula_headers = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}
    for row_idx, row in enumerate(rows, start=2):
        for key, value in row.items():
            col = headers.get(key)
            if col is None or key in formula_headers:
                continue
            if isinstance(value, (tuple, list, dict)):
                value = json.dumps(value, ensure_ascii=False)
            if isinstance(value, float) and not math.isfinite(value):
                value = None
            ws.cell(row_idx, col).value = value
        if row_idx > 2:
            for col in range(1, ws.max_column + 1):
                source, target = ws.cell(2, col), ws.cell(row_idx, col)
                if source.has_style:
                    target._style = source._style
                target.number_format = source.number_format
    output = Path(out_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def _json_default(value):
    if hasattr(value, "item"):
        return value.item()
    if hasattr(value, "tolist"):
        return value.tolist()
    raise TypeError(f"cannot serialize {type(value).__name__}")
